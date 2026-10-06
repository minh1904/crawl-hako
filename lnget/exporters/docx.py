from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from lnget.exporters import GetImage, normalize_image
from lnget.models import Chapter, Novel, Volume


def export(path: Path, novel: Novel, volume: Volume, chapters: list[Chapter],
           cover: bytes | None, get_image: GetImage) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    doc.core_properties.title = f"{novel.title} — {volume.title}"
    doc.core_properties.author = novel.author or novel.translator

    if cover:
        _image(doc, cover, width=5.0)
    h = doc.add_heading(novel.title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in (volume.title, f"Tác giả: {novel.author}" if novel.author else "",
                 f"Nhóm dịch: {novel.translator}" if novel.translator else ""):
        if line:
            doc.add_paragraph(line).alignment = WD_ALIGN_PARAGRAPH.CENTER

    for ch in chapters:
        doc.add_page_break()
        doc.add_heading(ch.title, level=2).alignment = WD_ALIGN_PARAGRAPH.CENTER
        for el in ch.elements:
            if el.type == "text":
                p = doc.add_paragraph(el.text)
                p.paragraph_format.first_line_indent = Inches(0.3)
                continue
            data = get_image(el.url)
            if not (data and _image(doc, data)):
                p = doc.add_paragraph("[Ảnh không tải được]")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.runs[0].italic = True

    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    doc.save(str(tmp))
    tmp.replace(path)


def _image(doc, data: bytes, width: float = 4.5) -> bool:
    norm = normalize_image(data, reencode=True)  # python-docx kén một số biến thể JPEG/PNG
    if not norm:
        return False
    w, h = norm[2]
    # ảnh dọc cao quá thì giới hạn theo chiều cao trang
    width = min(width, 8.0 * w / h) if h else width
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        p.add_run().add_picture(io.BytesIO(norm[0]), width=Inches(width))
    except Exception:
        p._element.getparent().remove(p._element)
        return False
    return True
